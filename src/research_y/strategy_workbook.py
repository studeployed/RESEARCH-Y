"""Render a frozen strategy/count snapshot to genuine Excel workbook bytes.

Internal, side-effect-free report renderer, NOT an audited file-export command.
No user-selected output is opened and no external workbook is read. The backend
may use temporary files internally. Persistent
export receipts, publication/recovery and user interface are separate gates.
"""
from __future__ import annotations

from io import BytesIO
import re
from typing import Any

from research_y.strategy_report import ReportIntegrityError, StrategyReport, canonical, digest

RENDER_METHOD = "strategy-count-xlsx/1"
ROW_LIMIT = 1_048_576
COLUMN_LIMIT = 16_384
TEXT_UNITS = 30_000  # Below Excel's 32,767-character cell limit, including UTF-16.
UNKNOWN = "Unknown"


class WorkbookLimitError(ValueError):
    """Output exceeds this renderer's supported size; nothing was published."""


class WorkbookTextError(ValueError):
    """Text cannot be represented losslessly in an XML workbook."""


def _text(value: str) -> None:
    # XML 1.0 permitted characters. Refuse rather than silently remove evidence.
    if any(not (c in "\t\n\r" or 0x20 <= ord(c) <= 0xD7FF
                or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF) for c in value):
        raise WorkbookTextError("Unsupported control character in report text")


def text_parts(value: str) -> tuple[str, ...]:
    """Lossless chunks, never split a Unicode scalar; order reconstructs text."""
    _text(value)
    result, part, units = [], [], 0
    for character in value:
        size = 2 if ord(character) > 0xFFFF else 1
        if units + size > TEXT_UNITS:
            result.append("".join(part))
            part, units = [], 0
        part.append(character)
        units += size
    result.append("".join(part))
    return tuple(result)


def _literal(cell, value: Any) -> None:
    if value is None:
        value = UNKNOWN
    if type(value) is int and abs(value) > 999_999_999_999_999:
        # Excel cannot represent arbitrary SQLite 64-bit integers exactly.
        value = str(value)
    if isinstance(value, str):
        _text(value)
        cell.value = value
        cell.data_type = "s"  # Explicit literal, including =,+,-,@ and error-like text.
        cell.number_format = "@"
    elif type(value) in (int, bool):
        cell.value = value
        if type(value) is int:
            cell.number_format = "#,##0"
    else:
        raise TypeError("Workbook cells require text, integers, booleans or explicit unknown")


def _methods(data, report):
    return [
        ("Report method", data["method"]), ("Workbook renderer", RENDER_METHOD),
        ("Software version", data["software_version"]),
        ("Project identifier", data["project"]["project_id"]), ("Project name", data["project"]["name"]),
        ("Run identifier", data["run"]["run_id"]), ("Strategy identifier", data["run"]["strategy_version_id"]),
        ("Strategy version", data["strategy"]["version_number"]),
        ("Approval time (UTC)", data["strategy"]["approved_at_utc"]),
        ("Captured at (UTC)", data["captured_at_utc"]), ("Run created at (UTC)", data["run"]["created_at_utc"]),
        ("Stored run status", data["run"]["status"]), ("Report status", "INCOMPLETE"),
        ("Scope", "Whole named run; no view filters applied; strategy/count report only"),
        ("Snapshot SHA-256", report.sha256), ("Count projection revision", data["counts"]["revision"]),
        ("Research theme", data["strategy"]["snapshot"]["theme"]),
        ("Boolean logic", "OR within and across approved groups"),
        ("Reported matches", "Source-reported total; Unknown is not zero. Multiple query totals are not added as a union."),
        ("Retrieved source records", "Unique provider identifiers; not a count of saved metadata or deduplicated publications."),
        ("Metadata source records", "Saved original source-record metadata with query membership."),
        ("Publication identities", "Current provisional identities, not verified deduplicated publication totals."),
        ("Group identity sum", "Arithmetic sum across groups; one identity can occur in several groups."),
        ("Global identity union", "Distinct active identities associated with this run; inactive retained identities excluded."),
        ("Verified unique publications", "Unknown: exhaustive candidate discovery/deduplication acceptance is outstanding."),
        ("Extracted abstracts", "Unknown: no completed abstract-extraction count is asserted."),
        ("Timestamp scope", "Executed Queries contains stored query timestamps, not independent verification of every request."),
        ("Query Partitions", "Stored PubMed count-plan nodes; internal split nodes are not extra independent searches or additive totals."),
        ("PubMed fields", "Connector targets All Fields literals; exact stored queries are authoritative. Accepted translations are not verified by this renderer."),
        ("Web of Science fields", "Starter Topic (TS): indexed title/abstract/keywords; narrower than web All Fields, not full article text."),
        ("Search limits", "Inspect exact stored queries and parameters. No inferred declaration of no limits."),
        ("Term provenance", "Approved atomic text, origin, rationale and enabled state retained; unavailable term ancestry is not invented."),
        ("PRISMA-S", "Preferred Reporting Items for Systematic Reviews and Meta-Analyses literature search extension: reporting evidence only; compliance not certified."),
        ("PRESS", "Peer Review of Electronic Search Strategies: peer-review completion not established by this workbook."),
        ("UTC", "Coordinated Universal Time; timestamps retained as exact sortable text."),
        ("JSON", "JavaScript Object Notation; exact structured report evidence."),
        ("SHA-256", "Secure Hash Algorithm 256-bit digest; integrity reference, not a digital signature or permission."),
        ("Literal cell policy", "All source/user text is explicitly stored as text. No formulas, macros or external hyperlinks are generated."),
        ("Long text", "A Continued reference identifies ordered Text Continuations rows; concatenate Part in order without separators."),
        ("Large integers", "Integers over 15 digits are exact text to prevent spreadsheet precision loss."),
        ("Snapshot scope", "Frozen present-day projection, not reconstruction of the run's original historical identity state."),
        ("Rights boundary", "No licensed titles, abstracts, provider payloads or record links are included. This is not a licence grant."),
        ("Unimplemented domains", "Master Results, Conference Abstracts, Society Guidelines and full Search Log are not exported; absence is not zero results."),
        ("Publication boundary", "Internal rendered workbook only. Audited file publication and recovery are not yet implemented."),
        ("Source status", "Numeric equality does not certify source completeness or establish a frozen provider snapshot."),
    ]


def _tables(data, report):
    counts = data["counts"]
    summary = []
    for group in counts["groups"]:
        for source in group["sources"]:
            summary.append((group["group_id"], group["label"], source["source_name"],
                source["reported_matches"], source["reported_count_status"], source["retrieved_source_records"],
                source["metadata_source_records"], source["identifiers_without_metadata"],
                source["numeric_reconciliation"], source["retrieved_minus_reported"]))
    yield "Search Summary", ("Group ID", "Group", "Source", "Reported matches", "Reported count status",
            "Retrieved identifiers", "Saved metadata records", "Identifiers without metadata", "Numeric reconciliation",
            "Retrieved minus reported"), summary
    yield "Source Totals", ("Source", "Unique source identifiers", "Saved metadata records", "Identifiers without metadata",
            "Queryless records", "Abstract availability (JSON)"), [
        (s["source_name"], s["source_unique_records"], s["metadata_source_records"], s["identifiers_without_metadata"],
         s["queryless_source_records"], canonical(s["abstract_status_counts"])) for s in counts["sources"]]
    pubrows = [(g["group_id"], g["label"], g["publication_identities"], g["group_unique_publications"],
                g["unmaterialised_retrievals"]) for g in counts["groups"]]
    yield "Publication Counts", ("Group ID", "Group", "Provisional identities", "Verified unique publications",
            "Unmaterialised retrievals"), pubrows
    yield "Run Totals", ("Named denominator or review state", "Value"), [
        *counts["totals"].items(), *(("review_" + k, v) for k, v in counts["reviews"].items())]
    terms = []
    for g in data["strategy"]["snapshot"]["groups"]:
        for order, term in enumerate(g["terms"]):
            terms.append((g["group_id"], g["label"], order + 1, term["text"], term["origin"],
                          term["approved"], term["enabled"], term["rationale"]))
    yield "Search Terms", ("Group ID", "Group", "Term order", "Approved atomic term", "Origin", "Approved",
                           "Enabled", "Rationale"), terms
    qcounts = {q["query_id"]: q for q in counts["queries"]}
    queries = []
    for q in data["queries"]:
        c = qcounts[q["query_id"]]
        queries.append((q["query_id"], q["group_id"], q["source_name"], q["exact_query"], canonical(q["parameters"]),
            q["executed_at_utc"], q["execution_evidence"], q["status"], q["reported_matches"],
            c["retrieved_source_records"], c["metadata_source_records"], "; ".join(c["warnings"])))
    yield "Executed Queries", ("Query ID", "Group ID", "Source", "Exact stored query", "Parameters (JSON)",
        "Stored timestamp (UTC)", "Execution evidence", "Stored query status", "Reported matches",
        "Retrieved identifiers", "Saved metadata", "Warnings"), queries
    yield "Query Partitions", ("Query ID", "Group ID", "Node", "Parent node", "Start", "End", "Exact subquery",
        "Reported node count", "Decision", "Depth", "Rationale", "Evidence scope"), [
        (p["query_id"], p["group_id"], p["index"], p["parent_index"], p["start"], p["end"], p["exact_subquery"],
         p["reported_matches"], p["decision"], p["depth"], p["rationale"], p["evidence_scope"]) for p in data["partitions"]]
    yield "Warnings", ("Scope", "Warning"), [("Whole run", w) for w in data["warnings"]]
    yield "Methods", ("Field or rule", "Definition or captured value"), _methods(data, report)
    parts = text_parts(report.payload_json)
    yield "Report Snapshot", ("Snapshot SHA-256", "Part", "Parts", "Canonical JSON chunk"), [
        (report.sha256, i + 1, len(parts), part) for i, part in enumerate(parts)]


def render_strategy_workbook(report: StrategyReport) -> bytes:
    """Render immutable evidence; dependency is optional outside this feature.

    Row/column overflows explicitly refuse output, never truncate. Long cells
    are losslessly continued. Audited file publication is deliberately absent.
    """
    if not isinstance(report, StrategyReport):
        raise TypeError("A captured StrategyReport is required")
    # Revalidate canonical evidence even if a caller bypassed frozen attributes.
    report = StrategyReport(report.payload_json)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("The optional Excel export dependencies are not installed") from None
    data = report.as_dict()
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "RESEARCH-Y"
    workbook.properties.title = "INCOMPLETE — strategy and count snapshot"
    workbook.properties.description = "Snapshot SHA-256: " + report.sha256
    continuations = []

    def table(name, headers, rows, *, continuation_table=False):
        if len(headers) > COLUMN_LIMIT or len(rows) + 4 > ROW_LIMIT:
            raise WorkbookLimitError("Workbook table exceeds row or column limit; no output published")
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        end = get_column_letter(len(headers))
        sheet.merge_cells(f"A1:{end}1")
        sheet.merge_cells(f"A2:{end}2")
        _literal(sheet["A1"], "RESEARCH-Y | " + name)
        _literal(sheet["A2"], "INCOMPLETE | Whole-run strategy/count snapshot | Provisional identities are not verified publications")
        sheet["A1"].font = Font(name="Calibri", size=17, bold=True, color="17365D")
        sheet["A2"].font = Font(name="Calibri", size=10, bold=True, color="9C2F00")
        sheet["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
        sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[2].height = 32
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(4, col)
            _literal(cell, header)
            cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(col)].width = 24
        sheet.row_dimensions[4].height = 34
        for row_number, row in enumerate(rows, 5):
            if len(row) != len(headers):
                raise ReportIntegrityError("Workbook row has unexpected columns")
            for col, value in enumerate(row, 1):
                cell = sheet.cell(row_number, col)
                if isinstance(value, str):
                    parts = text_parts(value)
                    if len(parts) > 1:
                        if continuation_table:
                            raise WorkbookLimitError("Continuation overflow")
                        key = f"{name}!{cell.coordinate}"
                        continuations.extend((key, i + 1, len(parts), part, digest(value))
                                             for i, part in enumerate(parts))
                        value = "[Continued in Text Continuations: " + key + "]"
                _literal(cell, value)
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if row_number % 2:
                    cell.fill = PatternFill("solid", fgColor="F2F5F9")
            sheet.row_dimensions[row_number].height = 42
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:{end}{max(4, len(rows) + 4)}"
        sheet.print_title_rows = "1:4"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        if name == "Methods":
            sheet.column_dimensions["A"].width = 34
            sheet.column_dimensions["B"].width = 90
        for header in ("Exact stored query", "Exact subquery", "Approved atomic term", "Rationale", "Canonical JSON chunk", "Text chunk"):
            if header in headers:
                sheet.column_dimensions[get_column_letter(headers.index(header) + 1)].width = 48
        # Only a fixed internal navigation link; no source-controlled URL targets.
        if name != "Methods":
            cell = sheet["A3"]
            _literal(cell, "Definitions and provenance")
            from openpyxl.worksheet.hyperlink import Hyperlink
            cell.hyperlink = Hyperlink(ref="A3", location="'Methods'!A1", display="Definitions and provenance")
            cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")

    try:
        for name, headers, rows in _tables(data, report):
            table(name, headers, rows)
        if continuations:
            table("Text Continuations", ("Original cell", "Part", "Parts", "Text chunk", "Full text SHA-256"),
                  continuations, continuation_table=True)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()

